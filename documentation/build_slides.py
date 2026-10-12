#!/usr/bin/env python3
"""Build the TDD PowerPoint for `cdc` from the snap.sh screenshots and the git history.

For every increment it makes:
  1 statement slide  (unit test, statement, expected results)
  5 screenshot slides (test code, test fails, code change, test passes, all tests pass)

usage (from the repo root; paths are resolved from the script's location, so any cwd works):
    uv run --with python-pptx python documentation/build_slides.py
    uv run --with python-pptx python documentation/build_slides.py -o cdc.pptx --authors "Alice, Bob"

Optional: documentation/statements.json overrides the auto-generated text, e.g.
    {
      "5":  {"statement": "Subtracting with one element must not lose it.",
             "expected":  ["raises IndexError", "stack still holds [3]"]},
      "14": {"note": "Already handled by the pop() implementation from inc 3."}
    }
"""
from __future__ import annotations

import argparse
import ast
import json
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- paths
# repo root, found from where this script lives (works from any current directory)
ROOT = Path(subprocess.check_output(["git", "rev-parse", "--show-toplevel"], text=True,
                                    cwd=Path(__file__).resolve().parent).strip())
SHOTS = ROOT / "documentation" / "shots"
OVERRIDES = ROOT / "documentation" / "statements.json"

# ---------------------------------------------------------------- style
DARK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x25, 0x63, 0xEB)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
MONO = "Consolas"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

# (file, title, caption, colour, short label for the tracker, commit phase)
STEPS = [
    ("1-test-code.png", "Unit test code", "New test added — diff of tests/", BLUE, "Test code", "RED"),
    ("2-test-fails.png", "Unit test fails (RED)", "Only the new test, run before implementing it", RED, "Fails", "RED"),
    ("3-code-change.png", "Code changes", "Highlighted diff of src/ (green = added, red = removed)", AMBER, "Code change", "GREEN"),
    ("4-test-passes.png", "Unit test passes (GREEN)", "Only the new test, run after the code change", GREEN, "Passes", "GREEN"),
    ("5-all-tests-pass.png", "All previous tests still pass", "Full regression run: uv run pytest", GREEN, "All pass", "GREEN"),
]

COMMIT_RE = re.compile(r"^inc(\d+)(?: (RED|GREEN))?: (\S+)(.*)$")


# ---------------------------------------------------------------- data
@dataclass
class Inc:
    n: int
    test: str
    note: str = ""
    commits: dict[str, str] = field(default_factory=dict)

    @property
    def folder(self) -> Path:
        return SHOTS / f"inc{self.n:02d}"

    @property
    def func(self) -> str:
        return self.test.split("::")[-1]

    @property
    def path(self) -> str:
        return self.test.split("::")[0]

    @property
    def passes_directly(self) -> bool:
        return "passes directly" in self.note

    def commit_for(self, phase: str) -> str | None:
        return self.commits.get(phase) or self.commits.get("SINGLE") or next(iter(self.commits.values()), None)


def read_history() -> list[Inc]:
    log = subprocess.check_output(["git", "log", "--reverse", "--format=%h%x09%s"], cwd=ROOT, text=True)
    incs: dict[int, Inc] = {}
    for line in log.splitlines():
        sha, _, subject = line.partition("\t")
        m = COMMIT_RE.match(subject)
        if not m:
            continue
        n = int(m[1])
        inc = incs.setdefault(n, Inc(n, m[3], m[4].strip(" ()")))
        inc.commits[m[2] or "SINGLE"] = sha
    # increments that have screenshots but no matching commit
    for d in SHOTS.glob("inc*"):
        if d.is_dir() and d.name[3:].isdigit():
            incs.setdefault(int(d.name[3:]), Inc(int(d.name[3:]), "(unknown test)"))
    return [incs[k] for k in sorted(incs)]


def find_test(inc: Inc) -> ast.FunctionDef | None:
    commit = inc.commit_for("GREEN")
    if not commit:
        return None
    try:
        src = subprocess.check_output(
            ["git", "show", f"{commit}:{inc.path}"], cwd=ROOT, text=True, stderr=subprocess.DEVNULL
        )
    except subprocess.CalledProcessError:
        return None
    for node in ast.walk(ast.parse(src)):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == inc.func:
            return node
    return None


def expectations(fn: ast.FunctionDef) -> list[str]:
    found: list[tuple[int, str]] = []
    for node in ast.walk(fn):
        if isinstance(node, ast.Assert):
            found.append((node.lineno, "assert " + ast.unparse(node.test)))
        elif isinstance(node, ast.With):
            for item in node.items:
                c = item.context_expr
                if isinstance(c, ast.Call) and ast.unparse(c.func).endswith("raises") and c.args:
                    found.append((node.lineno, f"raises {ast.unparse(c.args[0])}"))
    return [text for _, text in sorted(found)]


def humanize(func: str) -> str:
    words = func.removeprefix("test_").replace("_", " ").strip()
    return words[:1].upper() + words[1:]


def statement_for(inc: Inc, overrides: dict) -> tuple[str, list[str], str]:
    """Return (statement, expected results, note)."""
    o = overrides.get(str(inc.n), {})
    fn = find_test(inc)
    doc = ast.get_docstring(fn) if fn else None
    statement = o.get("statement") or (doc.strip() if doc else f"Checks that: {humanize(inc.func).lower()}.")
    expected = o.get("expected") or (expectations(fn) if fn else [])
    if not expected:
        expected = ["The test runs without raising an error."]
    note = o.get("note", "")
    if not note and inc.passes_directly:
        note = "This test passed directly: existing code already covered it, so no production change was needed."
    return statement, expected, note


# ---------------------------------------------------------------- drawing helpers
def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, lines, size=14, bold=False, color=DARK, font="Calibri",
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=4):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = color
    return tb


def picture(slide, path: Path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(str(path), Emu(x + (w - pw) // 2), Emu(y + (h - ph) // 2), Emu(pw), Emu(ph))


def header(slide, title, right, color):
    box(slide, 0, 0, SLIDE_W, Inches(0.95), color)
    text(slide, Inches(0.45), 0, Inches(8.4), Inches(0.95), title, size=24, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(8.9), 0, Inches(4.0), Inches(0.95), right, size=12, color=WHITE,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def footer(slide, left, right=""):
    text(slide, Inches(0.45), Inches(7.08), Inches(9), Inches(0.3), left, size=10, color=MUTED, font=MONO)
    text(slide, Inches(9.5), Inches(7.08), Inches(3.4), Inches(0.3), right, size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)


def tracker(slide, active: int | None, y=Inches(6.2)):
    """Row of the 5 TDD steps; `active` (0-4) is highlighted."""
    w, gap = Inches(2.3), Inches(0.18)
    x0 = (SLIDE_W - (5 * w + 4 * gap)) // 2
    for i, (*_, color, label, _phase) in enumerate(STEPS):
        on = active is None or i == active
        s = box(slide, x0 + i * (w + gap), y, w, Inches(0.5), color if on else LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
        tf = s.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = f"{i + 1}  {label}"
        r.font.size, r.font.bold = Pt(13), True
        r.font.color.rgb = WHITE if on else MUTED


# ---------------------------------------------------------------- slides
def title_slide(prs, title, subtitle, authors, n_incs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, SLIDE_W, SLIDE_H, DARK)
    box(s, Inches(0.8), Inches(3.55), Inches(1.4), Inches(0.08), GREEN)
    text(s, Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.4), title, size=48, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.BOTTOM)
    text(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.6), subtitle or
         f"Test-driven development in {n_incs} increments", size=22, color=RGBColor(0xD1, 0xD5, 0xDB))
    if authors:
        text(s, Inches(0.8), Inches(4.6), Inches(11.5), Inches(0.5), authors, size=16, color=MUTED)


def disclosure_slide(prs, n_incs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "How these slides were made", "Disclosure", DARK)

    cards = [
        (BLUE, "1  Captured during development",
         "While doing TDD, every RED and GREEN step ran our snap.sh script. It saved the "
         "screenshots and committed them with a structured message (incNN RED/GREEN: test id)."),
        (AMBER, "2  Built by a deterministic program",
         f"This deck was then generated automatically from the git history and those screenshots: "
         f"{n_incs} increments, no manual editing. The same repository always produces the same slides."),
        (GREEN, "3  Generator written with AI",
         "The script that builds the slides (build_slides.py) was written with the help of an AI "
         "assistant. It does not use AI when it runs: it only reads git, the test files and the images."),
    ]
    w, gap, top = Inches(3.95), Inches(0.29), Inches(1.45)
    for i, (color, title, body) in enumerate(cards):
        x = Inches(0.6) + i * (w + gap)
        box(s, x, top, w, Inches(3.0), LIGHT)
        box(s, x, top, w, Inches(0.09), color)
        text(s, x + Inches(0.25), top + Inches(0.3), w - Inches(0.5), Inches(0.5), title, size=17, bold=True,
             color=color)
        text(s, x + Inches(0.25), top + Inches(0.95), w - Inches(0.5), Inches(2.5), body, size=14)

    text(s, Inches(0.6), Inches(4.85), Inches(12.1), Inches(1.2),
         "This was only possible because of how we organised our TDD automation: each screenshot and "
         "commit was produced in a consistent, machine-readable way at the moment the work happened.",
         size=15, color=MUTED)

    box(s, Inches(0.6), Inches(5.65), Inches(12.13), Inches(0.85), LIGHT)
    box(s, Inches(0.6), Inches(5.65), Inches(0.09), Inches(0.85), DARK)
    text(s, Inches(0.9), Inches(5.65), Inches(11.7), Inches(0.85),
         "All of the calculator code, the unit tests and their statements, and the following of the TDD process itself "
         "were written by us, WITHOUT AI-generated code. AI was used only for writing the automation script and the slide-building script.",
         size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, "snap.sh  →  git history + documentation/shots/  →  build_slides.py  →  this deck")


def method_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "How each increment is documented", "", DARK)
    text(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(4.5), [
        "Every increment starts with a statement slide: the unit test, what it checks and the expected results.",
        "It is followed by five screenshots, generated automatically by snap.sh and committed with the code:",
        "1  the new unit test code (diff of tests/)",
        "2  the new test failing before any implementation — RED commit",
        "3  the production code changes, highlighted as a diff of src/",
        "4  the new test passing after the change — GREEN commit",
        "5  the whole suite (uv run pytest) still passing — no regressions",
        "The step bar at the bottom of each slide shows where you are in the cycle.",
    ], size=18, spacing=10)
    tracker(s, None)


def overview_slides(prs, incs):
    per = 14
    for start in range(0, len(incs), per):
        chunk = incs[start:start + per]
        s = prs.slides.add_slide(prs.slide_layouts[6])
        header(s, "Increments overview", f"{start + 1}–{start + len(chunk)} of {len(incs)}", DARK)
        rows = len(chunk) + 1
        tbl = s.shapes.add_table(rows, 3, Inches(0.45), Inches(1.2), Inches(12.43), Inches(0.38) * rows).table
        tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(0.9), Inches(9.3), Inches(2.23)
        for c, h in enumerate(["Inc", "Unit test", "File"]):
            tbl.cell(0, c).text = h
        for r, inc in enumerate(chunk, start=1):
            vals = [f"{inc.n:02d}", humanize(inc.func) + ("  (passed directly)" if inc.passes_directly else ""),
                    Path(inc.path).name]
            for c, v in enumerate(vals):
                tbl.cell(r, c).text = v
        for row in tbl.rows:
            for cell in row.cells:
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(12)


def statement_slide(prs, inc: Inc, statement, expected, note, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, f"Increment {inc.n:02d} — Unit test", f"{Path(inc.path).name}  ·  {inc.n} / {total}", DARK)

    text(s, Inches(0.6), Inches(1.25), Inches(12.1), Inches(0.7), humanize(inc.func), size=28, bold=True)
    text(s, Inches(0.6), Inches(1.95), Inches(12.1), Inches(0.4), inc.test, size=13, color=MUTED, font=MONO)

    # statement card
    box(s, Inches(0.6), Inches(2.55), Inches(5.6), Inches(3.35), LIGHT)
    text(s, Inches(0.85), Inches(2.75), Inches(5.1), Inches(0.4), "STATEMENT", size=12, bold=True, color=BLUE)
    text(s, Inches(0.85), Inches(3.15), Inches(5.1), Inches(2.6), statement, size=16)

    # expected-results card
    box(s, Inches(6.45), Inches(2.55), Inches(6.28), Inches(3.35), LIGHT)
    text(s, Inches(6.7), Inches(2.75), Inches(5.8), Inches(0.4), "EXPECTED RESULTS", size=12, bold=True, color=GREEN)
    shown = [e if len(e) <= 70 else e[:67] + "…" for e in expected[:8]]
    if len(expected) > 8:
        shown.append(f"… and {len(expected) - 8} more assertions")
    size = 13 if len(shown) <= 5 else 11
    text(s, Inches(6.7), Inches(3.15), Inches(5.8), Inches(2.6), [f"✓ {e}" for e in shown], size=size, font=MONO)

    if note:
        text(s, Inches(0.6), Inches(5.95), Inches(12.1), Inches(0.3), "ℹ " + note, size=12, color=AMBER, bold=True)
    tracker(s, None, y=Inches(6.4))
    commits = "  ".join(f"{k.lower()} {v}" for k, v in inc.commits.items() if k != "SINGLE") \
        or (f"commit {inc.commits['SINGLE']}" if "SINGLE" in inc.commits else "")
    footer(s, commits)


def step_slide(prs, inc: Inc, i: int, total):
    fname, title, caption, color, _label, phase = STEPS[i]
    if inc.passes_directly and i == 1:
        title, caption, color = "Unit test run before code change", "Passed directly — existing code already satisfied it", AMBER
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, f"Inc {inc.n:02d} · Step {i + 1}/5 — {title}", humanize(inc.func), color)
    text(s, Inches(0.45), Inches(1.05), Inches(12.4), Inches(0.35), caption, size=14, color=MUTED)

    img = inc.folder / fname
    x, y, w, h = Inches(0.45), Inches(1.45), Inches(12.43), Inches(4.6)
    if img.exists():
        picture(s, img, x, y, w, h)
    else:
        b = box(s, x, y, w, h, LIGHT)
        msg = ("No production code change was needed for this increment."
               if inc.passes_directly and i == 2 else f"No screenshot for this step ({fname} not found).")
        b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = b.text_frame.paragraphs[0].add_run()
        r.text, r.font.size, r.font.color.rgb = msg, Pt(18), MUTED

    tracker(s, i, y=Inches(6.3))
    sha = inc.commit_for(phase)
    footer(s, f"{img.relative_to(ROOT)}" + (f"   ·   commit {sha}" if sha else ""), f"{inc.n} / {total}")


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--output", default="documentation/cdc_tdd.pptx")
    ap.add_argument("--title", default="cdc — Complex Desk Calculator")
    ap.add_argument("--subtitle", default="")
    ap.add_argument("--authors", default="")
    args = ap.parse_args()

    overrides = json.loads(OVERRIDES.read_text()) if OVERRIDES.exists() else {}
    incs = read_history()
    if not incs:
        raise SystemExit("No incNN commits or screenshot folders found.")

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    title_slide(prs, args.title, args.subtitle, args.authors, len(incs))
    disclosure_slide(prs, len(incs))
    method_slide(prs)
    overview_slides(prs, incs)
    for inc in incs:
        statement, expected, note = statement_for(inc, overrides)
        statement_slide(prs, inc, statement, expected, note, len(incs))
        for i in range(len(STEPS)):
            step_slide(prs, inc, i, len(incs))

    out = ROOT / args.output
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    missing = sum(not (inc.folder / f).exists() for inc in incs for f, *_ in STEPS)
    print(f"Saved {out.relative_to(ROOT)}  ({len(prs.slides)} slides, {len(incs)} increments"
          + (f", {missing} screenshots missing" if missing else "") + ")")


if __name__ == "__main__":
    main()